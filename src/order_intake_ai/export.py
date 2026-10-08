"""ERP-ready outputs: one CSV + JSON of approved order lines, and an Excel workbook for the team."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ERP_COLUMNS = ["order_ref", "customer_id", "customer_name", "po_number", "order_date", "delivery_date",
               "line_no", "sku", "product", "unit", "quantity", "unit_price", "line_total"]


def erp_rows(report: dict) -> list[dict]:
    rows = []
    for o in report["orders"]:
        if o["status"] != "APPROVED":
            continue
        for ln in o["lines"]:
            m = ln["match"]
            rows.append({
                "order_ref": o["file"].rsplit(".", 1)[0], "customer_id": o["customer_id"], "customer_name": o["customer_name"],
                "po_number": o["po_number"], "order_date": o["order_date"], "delivery_date": o["delivery_date"] or "",
                "line_no": ln["line_no"], "sku": m["sku"], "product": m["name"], "unit": m["unit"],
                "quantity": ln["quantity"], "unit_price": ln["price"],
                "line_total": round(ln["price"] * ln["quantity"], 2) if ln["price"] is not None else None,
            })
    return rows


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for r in rows:
        ws.append(r)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
    ws.freeze_panes = "A2"


def write_outputs(report: dict, out_dir: Path) -> None:
    rows = erp_rows(report)
    with open(out_dir / "erp_import.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=ERP_COLUMNS)
        w.writeheader()
        w.writerows(rows)
    (out_dir / "erp_import.json").write_text(json.dumps(rows, indent=1, ensure_ascii=False))

    wb = Workbook()
    wb.remove(wb.active)
    orders = report["orders"]
    summary = report.get("summary") or {}
    lines_all = [ln for o in orders for ln in o["lines"]]
    auto = sum(ln["match"]["status"] == "AUTO" for ln in lines_all)
    kpis = [["Orders processed", len(orders)], ["Approved (ready for ERP)", sum(o["status"] == "APPROVED" for o in orders)],
            ["Sent to review", sum(o["status"] == "REVIEW" for o in orders)], ["Order lines", len(lines_all)],
            ["Lines matched automatically", auto]]
    if summary:
        kpis += [["Wrong automatic matches", summary["wrong_auto_matches"]],
                 ["Errors reaching the ERP", summary["errors_reaching_erp"]],
                 ["Planted problems caught", f"{summary['anomalies_caught']}/{summary['anomalies_total']}"]]
    _sheet(wb, "Summary", ["Metric", "Value"], kpis)
    _sheet(wb, "ERP import", ERP_COLUMNS, [[r[c] for c in ERP_COLUMNS] for r in rows])
    _sheet(wb, "Orders", ["File", "Status", "Customer", "PO", "Order date", "Delivery", "Lines", "Source", "Seconds", "Issues"],
           [[o["file"], o["status"], o["customer_name"] or o["sender"], o["po_number"], o["order_date"], o["delivery_date"],
             len(o["lines"]), o["route"], o["seconds"], "; ".join(x["message"] for x in o["issues"])] for o in orders])
    review = []
    for o in orders:
        for ln in o["lines"]:
            if ln["issues"]:
                cands = " | ".join(f"{c['sku']} {c['name']} ({c['score']:.2f})" for c in ln["match"]["candidates"])
                review.append([o["file"], o["customer_name"], ln["line_no"], ln["description"], ln["quantity"],
                               "; ".join(x["message"] for x in ln["issues"]), cands])
    _sheet(wb, "Review queue", ["File", "Customer", "Line", "Customer wrote", "Qty", "Why", "Top candidates"], review)
    _sheet(wb, "Lines", ["File", "Line", "Customer ref", "Customer wrote", "Qty", "Match", "Method", "SKU", "Product", "Score"],
           [[o["file"], ln["line_no"], ln["customer_ref"], ln["description"], ln["quantity"], ln["match"]["status"],
             ln["match"]["method"], ln["match"]["sku"], ln["match"]["name"], ln["match"]["score"]]
            for o in orders for ln in o["lines"]])
    if summary:
        acc = [[k, v] for k, v in summary.items() if not isinstance(v, (dict, list))]
        acc += [[f"header {k}", v] for k, v in summary["header_accuracy"].items()]
        acc += [[f"anomaly {k}", v] for k, v in summary["anomalies_by_type"].items()]
        _sheet(wb, "Accuracy", ["Metric", "Value"], acc)
    wb.save(out_dir / "orders.xlsx")
